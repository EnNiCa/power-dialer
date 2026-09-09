import csv
import io

from openpyxl import load_workbook

EXTENSIONES_SOPORTADAS = {'.csv', '.xlsx'}


class FormatoNoSoportado(Exception):
    pass


def parsear_clientes(fichero, nombre_fichero):
    """Lee un FileStorage de Flask y devuelve (filas, omitidas).

    `filas` es una lista de dicts {'nombre': str, 'telefono': str} ya
    limpiados. `omitidas` es una lista de números de fila (1 = primera fila
    de datos, sin contar la cabecera) que no tenían nombre o teléfono.
    """
    nombre_fichero = (nombre_fichero or '').lower()

    if nombre_fichero.endswith('.csv'):
        filas_crudas = _leer_csv(fichero)
    elif nombre_fichero.endswith('.xlsx'):
        filas_crudas = _leer_xlsx(fichero)
    else:
        raise FormatoNoSoportado('Formato no soportado, usa un fichero .csv o .xlsx')

    return _limpiar_filas(filas_crudas)


def _leer_csv(fichero):
    contenido = fichero.read().decode('utf-8-sig')

    try:
        dialecto = csv.Sniffer().sniff(contenido[:2048], delimiters=',;')
    except csv.Error:
        dialecto = csv.excel

    lector = csv.DictReader(io.StringIO(contenido), dialect=dialecto)
    return [_normalizar_cabeceras(fila) for fila in lector]


def _leer_xlsx(fichero):
    libro = load_workbook(fichero, read_only=True, data_only=True)
    hoja = libro.worksheets[0]
    filas = hoja.iter_rows(values_only=True)

    try:
        cabecera = [str(c or '').strip().lower() for c in next(filas)]
    except StopIteration:
        return []

    resultado = []
    for fila in filas:
        valores = {cabecera[i]: fila[i] for i in range(min(len(cabecera), len(fila)))}
        resultado.append(_normalizar_cabeceras(valores))
    return resultado


def _normalizar_cabeceras(fila):
    return {(clave or '').strip().lower(): valor for clave, valor in fila.items()}


def _texto(valor):
    """Convierte un valor de celda a texto, evitando que un teléfono leído
    como número en Excel se convierta en algo como '34611222333.0'."""
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _limpiar_filas(filas_crudas):
    filas = []
    omitidas = []

    for numero, fila in enumerate(filas_crudas, start=1):
        nombre = _texto(fila.get('nombre'))
        telefono = _texto(fila.get('telefono'))
        if not nombre or not telefono:
            omitidas.append(numero)
            continue
        filas.append({'nombre': nombre, 'telefono': telefono})

    return filas, omitidas
